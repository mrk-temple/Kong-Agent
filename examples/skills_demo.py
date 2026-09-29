"""Offline skills -> references -> XLSX creation -> structural verification."""
import asyncio
from pathlib import Path
import sys
from uuid import uuid4

from kong.contracts import Act, Action, Criterion, FinalResponse, Status
from kong.models.fake import ScriptedModel
from kong.runtime.loop import Runtime
from kong.storage import RunStore
from kong.tools.defaults import default_registry
from kong.workspace import Workspace


async def main():
    root = Path.cwd() / ".kong" / "demos" / uuid4().hex
    root.mkdir(parents=True)
    workspace = Workspace(root)
    registry = default_registry(workspace, allow_process=True)
    catalog = registry.get("skill_load").catalog
    health = catalog.load("spreadsheets")
    if health["missing"]:
        print("缺少必需依赖：" + ", ".join(health["missing"]))
        return 2
    script = """import csv, json
from pathlib import Path
from openpyxl import Workbook, load_workbook
with Path('source.csv').open(encoding='utf-8', newline='') as stream:
    records = list(csv.DictReader(stream))
book = Workbook()
sheet = book.active
sheet.title = 'Summary'
sheet.append(['ID', 'Amount'])
for row in records:
    sheet.append([row['ID'], int(row['Amount'])])
sheet.append(['Total', '=SUM(B2:B4)'])
book.save('report.xlsx')
check = load_workbook('report.xlsx', data_only=False)
assert check.active['A2'].value == '001'
assert check.active['B5'].value == '=SUM(B2:B4)'
total = sum(check.active.cell(row, 2).value for row in range(2, 5))
assert total == 16
Path('verification.json').write_text(json.dumps({'total': total, 'formula_preserved': True, 'recalculated': False}),encoding='utf-8')
print('Created and reopened report.xlsx; formula text preserved, recalculation not performed.')
"""
    decisions = [
        Act(actions=[Action(name="skill_load", arguments={"name": "spreadsheets"}),
                     Action(name="skill_read", arguments={"name": "spreadsheets", "resource": "references/workbooks.md"})]),
        Act(actions=[Action(name="write_file", arguments={"path": "source.csv", "content": "ID,Amount\n001,3\n002,5\n003,8\n"}),
                     Action(name="write_file", arguments={"path": "create_report.py", "content": script})]),
        Act(actions=[Action(name="process_exec", arguments={"argv": ["python", "create_report.py"]}),
                     Action(name="process_exec", arguments={"argv": ["python", str(catalog.get("spreadsheets").root / "scripts" / "inspect_file.py"), "report.xlsx", "--limit", "6"]}),
                     Action(name="read_file", arguments={"path": "verification.json"})]),
        FinalResponse(content="已生成 report.xlsx，保留前导零与公式；数据合计 16。公式尚未通过办公软件重算。", evidence_ids=["a3", "a4", "a5"]),
    ]
    runtime = Runtime.new("生成保留编号和公式的汇总表并验证", ScriptedModel(decisions), workspace, registry,
                          store=RunStore(root / ".kong" / "runs"))
    runtime.state.success_criteria = [Criterion(id="xlsx", description="workbook exists", kind="file_exists", path="report.xlsx"),
        Criterion(id="total", description="verified total", kind="file_contains", path="verification.json", contains='"total": 16')]
    print("固定决策离线演示：技能、参考文件和工具真实运行，不验证模型自动选技能的准确率。")
    state = await runtime.run()
    print(state.final_output or state.pending_question or state.stop_reason)
    print(f"工作区：{root.resolve()}")
    print(f"状态：{state.status}；进展证据不包含技能读取。")
    return 0 if state.status == Status.COMPLETED else 2


if __name__ == "__main__":
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    raise SystemExit(asyncio.run(main()))

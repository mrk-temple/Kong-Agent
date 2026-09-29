"""Read-only, bounded structural inspectors used by bundled office skills."""
import argparse
import csv
import json
from pathlib import Path
import sys
import zipfile

from kong.workspace import Workspace


def inspect_file(kind, path, offset=0, limit=20):
    path = Workspace(Path.cwd()).resolve(str(path))
    if not path.is_file() or path.stat().st_size > 32_000_000:
        raise ValueError("Input must be an existing workspace file no larger than 32 MB")
    if offset < 0 or not 1 <= limit <= 50:
        raise ValueError("offset >= 0 and 1 <= limit <= 50 required")
    if kind != "pdf" and path.suffix.lower() not in {".csv", ".tsv"}:
        with zipfile.ZipFile(path) as package:
            entries = package.infolist()
            if len(entries) > 10000 or sum(e.file_size for e in entries) > 64_000_000:
                raise ValueError("Office package exceeds inspection limits")
    result = {"file": str(path), "kind": kind, "offset": offset,
              "verification": "structural/text sample only; not rendering or formula recalculation"}
    if kind == "docx":
        from docx import Document
        doc = Document(path)
        result.update(paragraph_count=len(doc.paragraphs), table_count=len(doc.tables),
            paragraphs=[p.text[:500] for p in doc.paragraphs[offset:offset + limit]],
            tables=[[[c.text[:120] for c in row.cells[:8]] for row in t.rows[:8]] for t in doc.tables[:4]],
            truncated=True)
    elif kind == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("Encrypted PDF requires an authorized password workflow")
        pages = [{"page": i + 1, "text": (reader.pages[i].extract_text() or "")[:1500]}
                 for i in range(offset, min(len(reader.pages), offset + limit))]
        result.update(page_count=len(reader.pages), pages=pages, truncated=True,
                      empty_text_pages=[p["page"] for p in pages if not p["text"].strip()],
                      note="Empty extracted text may indicate a scan; OCR is not performed.")
    elif kind == "xlsx":
        if path.suffix.lower() in {".csv", ".tsv"}:
            with path.open(encoding="utf-8-sig", newline="") as stream:
                reader = csv.reader(stream, delimiter="\t" if path.suffix.lower() == ".tsv" else ",")
                rows = []
                for index, row in enumerate(reader):
                    if index < offset:
                        continue
                    if len(rows) == limit:
                        break
                    rows.append([str(v)[:200] for v in row[:20]])
            result.update(rows=rows, truncated=True)
        else:
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                result.update(sheet_count=len(workbook.sheetnames), sheets=[{
                    "name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column,
                    "sample": [[str(c.value)[:200] if c.value is not None else None for c in row]
                               for row in sheet.iter_rows(min_row=offset + 1, max_row=offset + limit,
                                                         max_col=min(sheet.max_column or 1, 20))]}
                    for sheet in list(workbook.worksheets)[:8]], truncated=True)
            finally:
                workbook.close()
    elif kind == "pptx":
        from pptx import Presentation
        deck = Presentation(path)
        slides = []
        for index in range(offset, min(len(deck.slides), offset + limit)):
            slide = deck.slides[index]
            texts = []
            for shape in list(slide.shapes)[:40]:
                if shape.has_text_frame:
                    texts.append(shape.text[:600])
                elif shape.has_table:
                    texts.extend(" | ".join(c.text[:100] for c in list(row.cells)[:8])
                                 for row in list(shape.table.rows)[:8])
            slides.append({"slide": index + 1, "shape_count": len(slide.shapes), "text": texts[:40]})
        result.update(slide_count=len(deck.slides), slides=slides, truncated=True)
    else:
        raise ValueError("Unknown inspection format")
    return result


def main(kind):
    parser = argparse.ArgumentParser(description="Inspect a workspace office file without modifying it")
    parser.add_argument("path")
    parser.add_argument("--offset", type=int, default=0)
    parser.add_argument("--limit", type=int, default=20)
    args = parser.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    try:
        print(json.dumps(inspect_file(kind, args.path, args.offset, args.limit), ensure_ascii=False))
        return 0
    except Exception as exc:
        print(json.dumps({"error": type(exc).__name__, "detail": str(exc)[:500]}, ensure_ascii=False))
        return 2
